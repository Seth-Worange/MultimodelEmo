"""Compare untouched human word boundaries with old and local-crop MFA."""

from __future__ import annotations

import argparse
import csv
import math
import shutil
import statistics
from pathlib import Path
from typing import Any

from .alignment import normalize_word
from .io_utils import sha256_file, write_csv, write_json


REFERENCE_FIELDS = (
    "sample_id", "word_index", "original_word", "reference_start_s",
    "reference_end_s", "reviewer_id", "boundary_status", "source_sheet", "source_row",
)
DETAIL_FIELDS = (
    "sample_id", "word_index", "original_word", "boundary_status", "reviewer_id",
    "manual_start", "manual_end", "auto_start", "auto_end", "start_error",
    "end_error", "max_error", "old_auto_start", "old_auto_end",
    "old_start_error", "old_end_error", "old_max_error", "comparison_status",
)


def _read_alignment(path: Path) -> dict[int, dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    output = {int(row["word_index"]): row for row in rows}
    if len(output) != len(rows):
        raise ValueError(f"Duplicate word index in {path}")
    return output


def _read_reference_xlsx(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    header = [str(value).strip().lower() if value is not None else "" for value in next(iterator)]
    required = {"sample_id", "word_index", "reference_start_s", "reference_end_s", "reviewer_id"}
    if not required.issubset(header):
        raise ValueError(f"Manual file missing columns: {sorted(required - set(header))}")
    positions = {name: header.index(name) for name in required}
    original_position = header.index("original words") if "original words" in header else None
    output = []
    for source_row, values in enumerate(iterator, 2):
        if not any(value is not None for value in values):
            continue
        start = values[positions["reference_start_s"]]
        end = values[positions["reference_end_s"]]
        valid = isinstance(start, (int, float)) and isinstance(end, (int, float))
        if valid and not (math.isfinite(start) and math.isfinite(end) and 0 <= start < end):
            valid = False
        output.append({
            "sample_id": str(values[positions["sample_id"]]),
            "word_index": int(values[positions["word_index"]]),
            "original_word": str(values[original_position]) if original_position is not None else "",
            "reference_start_s": start if start is not None else "",
            "reference_end_s": end if end is not None else "",
            "reviewer_id": str(values[positions["reviewer_id"]]) if values[positions["reviewer_id"]] is not None else "",
            "boundary_status": "provided_numeric_boundaries" if valid else "boundary_uncertain_or_missing",
            "source_sheet": sheet.title,
            "source_row": source_row,
        })
    workbook.close()
    if len({(row["sample_id"], row["word_index"]) for row in output}) != len(output):
        raise ValueError("Duplicate reference word index")
    return output


def _variant_metrics(rows: list[dict[str, Any]], prefix: str) -> dict[str, Any]:
    if prefix == "new":
        start_key, end_key, max_key = "start_error", "end_error", "max_error"
    else:
        start_key, end_key, max_key = "old_start_error", "old_end_error", "old_max_error"
    valid = [row for row in rows if row[start_key] != "" and row[end_key] != ""]
    starts = [float(row[start_key]) for row in valid]
    ends = [float(row[end_key]) for row in valid]
    boundaries = starts + ends
    return {
        "reference_word_count": sum(row["boundary_status"] == "provided_numeric_boundaries" for row in rows),
        "compared_word_count": len(valid),
        "start_mae_s": statistics.mean(starts) if starts else None,
        "end_mae_s": statistics.mean(ends) if ends else None,
        "mean_boundary_error_s": statistics.mean(boundaries) if boundaries else None,
        "median_boundary_error_s": statistics.median(boundaries) if boundaries else None,
        "max_boundary_error_s": max(boundaries) if boundaries else None,
        "pass_rate_50ms": sum(float(row[max_key]) <= .05 for row in valid) / len(valid) if valid else None,
        "pass_rate_100ms": sum(float(row[max_key]) <= .10 for row in valid) / len(valid) if valid else None,
        "pass_rate_200ms": sum(float(row[max_key]) <= .20 for row in valid) / len(valid) if valid else None,
        "pass_rule": "both word boundaries within tolerance (max absolute boundary error)",
    }


def validate_reference(
    manual_xlsx: Path, old_alignment_csv: Path, new_alignment_csv: Path,
    output_dir: Path,
) -> dict[str, Any]:
    if not manual_xlsx.is_file():
        raise FileNotFoundError(manual_xlsx)
    output_dir.mkdir(parents=True, exist_ok=True)
    source_copy = output_dir / manual_xlsx.name
    if not source_copy.exists():
        shutil.copy2(manual_xlsx, source_copy)
    elif sha256_file(source_copy) != sha256_file(manual_xlsx):
        raise ValueError("Existing manual reference copy differs; refusing overwrite")
    old = _read_alignment(old_alignment_csv)
    new = _read_alignment(new_alignment_csv)
    reference = _read_reference_xlsx(manual_xlsx)
    write_csv(output_dir / "manual_alignment_reference.csv", reference, REFERENCE_FIELDS)
    detail = []
    for row in reference:
        index = row["word_index"]
        if index not in old or index not in new:
            raise ValueError(f"Manual word {index} absent in automatic alignment")
        old_row, new_row = old[index], new[index]
        for variant in (old_row, new_row):
            if variant["sample_id"] != row["sample_id"]:
                raise ValueError(f"Sample mismatch at word {index}")
            if normalize_word(variant["original_word"]) != normalize_word(row["original_word"]):
                raise ValueError(f"Word mismatch at position {index}: {row['original_word']} / {variant['original_word']}")
        valid_reference = row["boundary_status"] == "provided_numeric_boundaries"

        def errors(variant: dict[str, str]) -> tuple[float | str, float | str, float | str, float | str, float | str]:
            if not valid_reference or variant["alignment_mask"] != "1":
                return "", "", "", "", ""
            start = float(variant["start_s"])
            end = float(variant["end_s"])
            start_error = abs(start - float(row["reference_start_s"]))
            end_error = abs(end - float(row["reference_end_s"]))
            return start, end, start_error, end_error, max(start_error, end_error)

        new_start, new_end, start_error, end_error, max_error = errors(new_row)
        old_start, old_end, old_start_error, old_end_error, old_max_error = errors(old_row)
        detail.append({
            "sample_id": row["sample_id"], "word_index": index,
            "original_word": row["original_word"],
            "boundary_status": row["boundary_status"], "reviewer_id": row["reviewer_id"],
            "manual_start": row["reference_start_s"], "manual_end": row["reference_end_s"],
            "auto_start": new_start, "auto_end": new_end,
            "start_error": start_error, "end_error": end_error, "max_error": max_error,
            "old_auto_start": old_start, "old_auto_end": old_end,
            "old_start_error": old_start_error, "old_end_error": old_end_error,
            "old_max_error": old_max_error,
            "comparison_status": "both_available" if max_error != "" and old_max_error != "" else "missing_boundary",
        })
    write_csv(output_dir / "manual_alignment_detail.csv", detail, DETAIL_FIELDS)
    metrics = {
        "sample_ids": sorted({row["sample_id"] for row in reference}),
        "manual_source_path": str(manual_xlsx.resolve()),
        "manual_source_sha256": sha256_file(manual_xlsx),
        "manual_source_reviewer_id_blank_count": sum(not row["reviewer_id"] for row in reference),
        "reference_status": "user_supplied_numeric_boundaries; reviewer_id not independently verified",
        "old_whole_audio_mfa": _variant_metrics(detail, "old"),
        "new_local_crop_mfa": _variant_metrics(detail, "new"),
        "note": "Boundary evaluation is restricted to these human-provided word rows, not the 100-sample automatic QA.",
    }
    write_json(output_dir / "manual_alignment_metrics.json", metrics)
    return metrics


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manual-xlsx", type=Path, required=True)
    parser.add_argument("--old-alignment-csv", type=Path, required=True)
    parser.add_argument("--new-alignment-csv", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    print(validate_reference(
        args.manual_xlsx, args.old_alignment_csv,
        args.new_alignment_csv, args.output_dir,
    ))


if __name__ == "__main__":
    main()
