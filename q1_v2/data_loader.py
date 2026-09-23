"""Read and audit the immutable Attachment 1 sample set."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .io_utils import write_csv, write_json


REQUIRED_COLUMNS = ("video_id", "clip_id", "text", "label", "annotation")


@dataclass(frozen=True)
class OriginalWord:
    index: int
    text: str
    char_start: int
    char_end: int


@dataclass
class SampleRecord:
    sample_id: str
    video_id: str
    clip_id: str
    text: str
    video_path: Path
    source_row: int
    original_words: list[OriginalWord]
    # Audit-only values. Feature modules never receive these two fields.
    label_audit: str = ""
    annotation_audit: str = ""
    read_status: str = "pending"
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def word_count(self) -> int:
        return len(self.original_words)

    def manifest_row(self) -> dict[str, Any]:
        return {
            "sample_id": self.sample_id,
            "video_id": self.video_id,
            "clip_id": self.clip_id,
            "source_row": self.source_row,
            "video_path": str(self.video_path),
            "video_exists": int(self.video_path.is_file()),
            "text": self.text,
            "word_count": self.word_count,
            "label_audit": self.label_audit,
            "annotation_audit": self.annotation_audit,
            "read_status": self.read_status,
            "errors": " | ".join(self.errors),
            "warnings": " | ".join(self.warnings),
        }


MANIFEST_FIELDS = (
    "sample_id", "video_id", "clip_id", "source_row", "video_path",
    "video_exists", "text", "word_count", "label_audit",
    "annotation_audit", "read_status", "errors", "warnings",
)


def canonical_id(value: Any) -> str:
    """Canonicalize Excel identifiers without altering meaningful strings."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def make_sample_id(video_id: str, clip_id: str) -> str:
    # MOSEI video IDs and numeric clip IDs do not contain this separator.
    return f"{video_id}__{clip_id}"


def split_original_words(text: str) -> list[OriginalWord]:
    """Preserve the exact non-whitespace tokens and their source offsets."""
    return [
        OriginalWord(index=i, text=match.group(0), char_start=match.start(), char_end=match.end())
        for i, match in enumerate(re.finditer(r"\S+", text))
    ]


def resolve_label_path(data_root: Path, labels: Path | None) -> Path:
    if labels is not None:
        path = labels.resolve()
        if not path.is_file():
            raise FileNotFoundError(f"Label workbook does not exist: {path}")
        return path
    candidates = sorted(data_root.rglob("label-100.xlsx"))
    if len(candidates) != 1:
        raise FileNotFoundError(
            f"Expected exactly one label-100.xlsx below {data_root}, found {len(candidates)}; "
            "pass --labels explicitly"
        )
    return candidates[0].resolve()


def _probe_video(path: Path) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    if not path.is_file():
        return ["missing_video_file"], warnings
    try:
        import av

        with av.open(str(path)) as container:
            if not container.streams.video:
                errors.append("missing_video_stream")
            else:
                try:
                    next(container.decode(container.streams.video[0]))
                except StopIteration:
                    errors.append("video_stream_has_no_decodable_frames")
            if not container.streams.audio:
                errors.append("missing_audio_stream")
    except ImportError:
        warnings.append("pyav_unavailable_decode_not_checked")
    except Exception as exc:  # The exception type/message is part of the audit trail.
        errors.append(f"video_decode_error:{type(exc).__name__}:{exc}")
    return errors, warnings


def load_samples(
    data_root: Path,
    labels: Path | None = None,
    *,
    expected_count: int = 100,
    probe_decode: bool = True,
    output_dir: Path | None = None,
) -> tuple[list[SampleRecord], dict[str, Any]]:
    """Read the workbook and verify one-to-one workbook/video correspondence."""
    from openpyxl import load_workbook

    data_root = data_root.resolve()
    if not data_root.is_dir():
        raise NotADirectoryError(f"Attachment root does not exist: {data_root}")
    label_path = resolve_label_path(data_root, labels)
    workbook = load_workbook(label_path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header_values = next(rows)
    except StopIteration as exc:
        raise ValueError(f"Empty workbook: {label_path}") from exc
    headers = [str(value).strip() if value is not None else "" for value in header_values]
    missing_columns = [name for name in REQUIRED_COLUMNS if name not in headers]
    if missing_columns:
        raise ValueError(f"Workbook missing columns: {missing_columns}")
    positions = {name: headers.index(name) for name in REQUIRED_COLUMNS}

    records: list[SampleRecord] = []
    seen: dict[str, int] = {}
    for excel_row, values in enumerate(rows, start=2):
        if not any(value is not None for value in values):
            continue
        video_id = canonical_id(values[positions["video_id"]])
        clip_id = canonical_id(values[positions["clip_id"]])
        text_value = values[positions["text"]]
        text = "" if text_value is None else str(text_value)
        sample_id = make_sample_id(video_id, clip_id)
        video_path = data_root / video_id / f"{clip_id}.mp4"
        record = SampleRecord(
            sample_id=sample_id,
            video_id=video_id,
            clip_id=clip_id,
            text=text,
            video_path=video_path,
            source_row=excel_row,
            original_words=split_original_words(text),
            label_audit=canonical_id(values[positions["label"]]),
            annotation_audit=canonical_id(values[positions["annotation"]]),
        )
        if not video_id:
            record.errors.append("empty_video_id")
        if not clip_id:
            record.errors.append("empty_clip_id")
        if not text.strip():
            record.errors.append("empty_transcript")
        if sample_id in seen:
            record.errors.append(f"duplicate_sample_id:first_row={seen[sample_id]}")
        else:
            seen[sample_id] = excel_row
        if probe_decode:
            errors, warnings = _probe_video(video_path)
            record.errors.extend(errors)
            record.warnings.extend(warnings)
        elif not video_path.is_file():
            record.errors.append("missing_video_file")
        record.read_status = "ok" if not record.errors else "invalid"
        records.append(record)
    workbook.close()

    referenced = {record.video_path.resolve() for record in records}
    discovered = {path.resolve() for path in data_root.rglob("*.mp4")}
    extra_videos = sorted(str(path) for path in discovered - referenced)
    missing_videos = sorted(str(path) for path in referenced - discovered)
    summary = {
        "label_path": str(label_path),
        "data_root": str(data_root),
        "row_count": len(records),
        "expected_count": expected_count,
        "count_matches_expected": len(records) == expected_count,
        "unique_sample_count": len(seen),
        "video_file_count": len(discovered),
        "valid_sample_count": sum(record.read_status == "ok" for record in records),
        "invalid_sample_count": sum(record.read_status != "ok" for record in records),
        "missing_videos": missing_videos,
        "extra_videos": extra_videos,
    }
    if len(records) != expected_count:
        summary["dataset_error"] = f"expected_{expected_count}_rows_found_{len(records)}"

    if output_dir is not None:
        output_dir.mkdir(parents=True, exist_ok=True)
        write_csv(output_dir / "input_manifest.csv", (item.manifest_row() for item in records), MANIFEST_FIELDS)
        write_json(output_dir / "data_loading_summary.json", summary)
        log_lines = [
            f"rows={len(records)} expected={expected_count} unique={len(seen)} videos={len(discovered)}",
            f"valid={summary['valid_sample_count']} invalid={summary['invalid_sample_count']}",
        ]
        for record in records:
            if record.errors or record.warnings:
                log_lines.append(
                    f"{record.sample_id}\tstatus={record.read_status}\t"
                    f"errors={' | '.join(record.errors)}\twarnings={' | '.join(record.warnings)}"
                )
        (output_dir / "data_loading.log").write_text("\n".join(log_lines) + "\n", encoding="utf-8")
    return records, summary
