"""提取100段视频的词级音视频特征及对齐信息。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import subprocess
import sys
from difflib import SequenceMatcher
from pathlib import Path

import numpy as np
import torch
from openpyxl import load_workbook

from utils.config import parse_config_args

SAMPLE_RATE = 16000
N_MELS = 40
N_FFT = 512
FRAME = 400
HOP = 160
# 鼻尖仅作原点，输出肩、肘、腕和髋部姿态。
FACE_POINTS = (10, 151, 9, 168, 2, 33, 133, 159, 145, 362, 263, 386, 374, 61, 291, 13)
BODY_POINTS = (11, 12, 13, 14, 15, 16, 23, 24)
FACE_DIM = len(FACE_POINTS) * 2 + 1
BODY_DIM = len(BODY_POINTS) * 4 + 1
VISION_FRAME_DIM = FACE_DIM + BODY_DIM
REVIEW_ALIGNMENT_SCORE = 0.3  # 低于阈值的样本需人工复核。
PROSODY_FIELDS = ("duration_s", "pause_before_s", "pause_after_s", "local_words_per_s",
                  "relative_f0_mean_st", "f0_range_st", "f0_slope_st_per_s", "voiced_fraction",
                  "log_rms_mean", "log_rms_range", "log_rms_slope_per_s", "zcr_mean")


def configure_cache() -> None:
    cache = Path(__file__).resolve().parent.parent / "cache"
    for name, suffix in (("HF_HOME", "huggingface"), ("TORCH_HOME", "torch"),
                         ("NLTK_DATA", "nltk_data")):
        os.environ.setdefault(name, str(cache / suffix))


def hertz_to_mel(hz: np.ndarray | float) -> np.ndarray:
    return 2595.0 * np.log10(1.0 + np.asarray(hz) / 700.0)


def mel_bank() -> np.ndarray:
    mel_points = np.linspace(hertz_to_mel(0), hertz_to_mel(SAMPLE_RATE / 2), N_MELS + 2)
    hz_points = 700.0 * (10 ** (mel_points / 2595.0) - 1.0)
    bins = np.floor((N_FFT + 1) * hz_points / SAMPLE_RATE).astype(int)
    bank = np.zeros((N_MELS, N_FFT // 2 + 1), dtype=np.float32)
    for m in range(N_MELS):
        left, center, right = bins[m:m + 3]
        for k in range(left, center):
            bank[m, k] = (k - left) / max(1, center - left)
        for k in range(center, right):
            bank[m, k] = (right - k) / max(1, right - center)
    return bank


def decode_audio(video: Path) -> np.ndarray:
    command = ["ffmpeg", "-hide_banner", "-loglevel", "error", "-i", str(video), "-vn",
               "-ac", "1", "-ar", str(SAMPLE_RATE), "-f", "s16le", "pipe:1"]
    result = subprocess.run(command, check=True, capture_output=True)
    audio = np.frombuffer(result.stdout, dtype="<i2").astype(np.float32) / 32768.0
    if not len(audio):
        raise ValueError(f"No audio decoded from {video}")
    return audio


def audio_frames(audio: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    bank = mel_bank()
    starts = np.arange(0, max(1, len(audio) - FRAME + 1), HOP)
    features = []
    window = np.hanning(FRAME).astype(np.float32)
    min_lag, max_lag = int(SAMPLE_RATE / 400), int(SAMPLE_RATE / 60)
    for start in starts:
        frame = audio[start:start + FRAME]
        if len(frame) < FRAME:
            frame = np.pad(frame, (0, FRAME - len(frame)))
        centered = frame - frame.mean()
        spectrum = np.abs(np.fft.rfft(centered * window, n=N_FFT)) ** 2
        mel = np.log(np.maximum(bank @ spectrum.astype(np.float32), 1e-10))
        rms = float(np.sqrt(np.mean(frame ** 2) + 1e-12))
        zcr = float(np.mean(np.signbit(frame[1:]) != np.signbit(frame[:-1])))
        corr_fft = np.fft.rfft(centered, n=1024)
        corr = np.fft.irfft(corr_fft * np.conj(corr_fft), n=1024)[:FRAME]
        upper = min(max_lag, len(corr) - 1)
        lag = min_lag + int(np.argmax(corr[min_lag:upper + 1])) if upper >= min_lag else 0
        pitch = SAMPLE_RATE / lag if lag and corr[lag] > 0.25 * max(corr[0], 1e-8) else 0.0
        features.append(np.r_[mel, np.log(rms + 1e-8), zcr, np.log(pitch + 1.0)].astype(np.float32))
    times = (starts + FRAME / 2) / SAMPLE_RATE
    return times.astype(np.float32), np.asarray(features, dtype=np.float32)


def align_words(transcript: str, audio: np.ndarray, align_model, metadata, whisperx, device: str) -> list[dict]:
    segments = [{"start": 0.0, "end": len(audio) / SAMPLE_RATE, "text": transcript}]
    aligned = whisperx.align(segments, align_model, metadata, audio, device, return_char_alignments=False)
    return [word for segment in aligned["segments"] for word in segment.get("words", [])
            if word.get("start") is not None and word.get("end") is not None]


def text_features(words: list[str], tokenizer, text_model, device: str) -> np.ndarray:
    encoded = tokenizer(words, is_split_into_words=True, return_tensors="pt", truncation=True,
                        max_length=512, padding=False)
    word_ids = encoded.word_ids(batch_index=0)
    encoded = {key: value.to(device) for key, value in encoded.items() if key != "offset_mapping"}
    with torch.inference_mode():
        hidden = text_model(**encoded).last_hidden_state[0].cpu().numpy()
    result = np.zeros((len(words), hidden.shape[-1]), dtype=np.float32)
    counts = np.zeros(len(words), dtype=np.int32)
    for token_index, word_index in enumerate(word_ids):
        if word_index is not None and word_index < len(words):
            result[word_index] += hidden[token_index]
            counts[word_index] += 1
    valid = counts > 0
    result[valid] /= counts[valid, None]
    return result


def visual_frames(video: Path, face_model: Path, pose_model: Path) -> tuple[np.ndarray, np.ndarray]:
    import cv2
    import mediapipe as mp

    capture = cv2.VideoCapture(str(video))
    if not capture.isOpened():
        capture.release()
        raise ValueError(f"Could not open video {video}")
    fps = capture.get(cv2.CAP_PROP_FPS)
    fps = fps if np.isfinite(fps) and fps > 0 else 30.0
    step = max(1, round(fps / 5.0))
    times, features = [], []
    face_landmarker = pose_landmarker = None
    try:
        face_options = mp.tasks.vision.FaceLandmarkerOptions(
            base_options=mp.tasks.BaseOptions(model_asset_buffer=face_model.read_bytes()),
            running_mode=mp.tasks.vision.RunningMode.IMAGE,
            num_faces=1,
        )
        pose_options = mp.tasks.vision.PoseLandmarkerOptions(
            base_options=mp.tasks.BaseOptions(model_asset_buffer=pose_model.read_bytes()),
            running_mode=mp.tasks.vision.RunningMode.IMAGE,
            num_poses=1,
        )
        face_landmarker = mp.tasks.vision.FaceLandmarker.create_from_options(face_options)
        pose_landmarker = mp.tasks.vision.PoseLandmarker.create_from_options(pose_options)
        frame_index = 0
        while True:
            ok, bgr = capture.read()
            if not ok:
                break
            if frame_index % step == 0:
                rgb = cv2.cvtColor(bgr, cv2.COLOR_BGR2RGB)
                image = mp.Image(image_format=mp.ImageFormat.SRGB, data=rgb)
                face_result = face_landmarker.detect(image)
                pose_result = pose_landmarker.detect(image)
                time_sec = frame_index / fps
                face = np.zeros(FACE_DIM, dtype=np.float32)
                if face_result.face_landmarks:
                    points = face_result.face_landmarks[0]
                    # 以鼻尖为原点，以双眼间距消除尺度差异。
                    nose = points[1]
                    eye_distance = max(np.hypot(points[33].x - points[263].x,
                                                points[33].y - points[263].y), 1e-4)
                    coords = []
                    for point_id in FACE_POINTS:
                        point = points[point_id]
                        coords.extend(((point.x - nose.x) / eye_distance,
                                       (point.y - nose.y) / eye_distance))
                    face[:] = np.r_[coords, 1.0]

                body = np.zeros(BODY_DIM, dtype=np.float32)
                if pose_result.pose_landmarks:
                    points = pose_result.pose_landmarks[0]
                    # 以双肩中心和肩宽归一化上半身坐标。
                    left, right = points[11], points[12]
                    width = max(np.hypot((left.x - right.x) * bgr.shape[1],
                                         (left.y - right.y) * bgr.shape[0]), 1e-4)
                    center_x = (left.x + right.x) * bgr.shape[1] / 2
                    center_y = (left.y + right.y) * bgr.shape[0] / 2
                    coords = []
                    for point_id in BODY_POINTS:
                        point = points[point_id]
                        coords.extend(((point.x * bgr.shape[1] - center_x) / width,
                                       (point.y * bgr.shape[0] - center_y) / width,
                                       point.z * bgr.shape[1] / width,
                                       float(point.visibility or 0.0)))
                    body[:] = np.r_[coords, 1.0]
                features.append(np.r_[face, body])
                times.append(time_sec)
            frame_index += 1
    finally:
        if face_landmarker is not None:
            face_landmarker.close()
        if pose_landmarker is not None:
            pose_landmarker.close()
        capture.release()
    if not features:
        raise ValueError(f"No video frames decoded from {video}")
    return np.asarray(times, dtype=np.float32), np.asarray(features, dtype=np.float32)


def pool_intervals(times: np.ndarray, features: np.ndarray, words: list[dict],
                   nearest_tolerance: float = 0.0) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    width = features.shape[1] * 2
    pooled = np.zeros((len(words), width), dtype=np.float32)
    valid = np.zeros(len(words), dtype=np.uint8)
    nearest = np.zeros(len(words), dtype=np.uint8)
    for i, word in enumerate(words):
        start, end = float(word["start"]), float(word["end"])
        selected = (times >= start) & (times < end)
        if not selected.any() and nearest_tolerance and len(times):
            gaps = np.maximum(start - times, 0) + np.maximum(times - end, 0)
            closest = int(np.argmin(gaps))
            if gaps[closest] <= nearest_tolerance:
                selected[closest] = True
                nearest[i] = 1
        if selected.any():
            values = features[selected]
            pooled[i] = np.r_[values.mean(axis=0), values.std(axis=0)]
            valid[i] = 1
    return pooled, valid, nearest


def prosody_features(times: np.ndarray, features: np.ndarray, words: list[dict],
                     utterance_duration: float) -> np.ndarray:
    """汇总词时长、音高、浊音和能量变化。"""
    pitch = np.expm1(features[:, 42]).clip(min=0.0)
    voiced = pitch > 0
    reference_pitch = float(np.median(pitch[voiced])) if voiced.any() else 1.0
    relative_f0 = np.zeros_like(pitch)
    relative_f0[voiced] = 12.0 * np.log2(pitch[voiced] / max(reference_pitch, 1e-6))
    energy, zcr = features[:, 40], features[:, 41]
    starts = np.asarray([float(word["start"]) for word in words])
    ends = np.asarray([float(word["end"]) for word in words])
    centers = (starts + ends) / 2
    result = np.zeros((len(words), len(PROSODY_FIELDS)), dtype=np.float32)

    def slope(x: np.ndarray, y: np.ndarray) -> float:
        if len(x) < 2:
            return 0.0
        centered = x - x.mean()
        denominator = float(centered @ centered)
        return float(centered @ (y - y.mean()) / denominator) if denominator > 1e-8 else 0.0

    for i, word in enumerate(words):
        start, end = starts[i], ends[i]
        selected = (times >= start) & (times < end)
        voiced_word = selected & voiced
        local_start, local_end = max(0.0, centers[i] - 0.5), min(utterance_duration, centers[i] + 0.5)
        local_duration = max(local_end - local_start, 1e-6)
        local_rate = np.sum((centers >= local_start) & (centers < local_end)) / local_duration
        before = max(0.0, start - ends[i - 1]) if i else 0.0
        after = max(0.0, starts[i + 1] - end) if i + 1 < len(words) else 0.0
        f0 = relative_f0[voiced_word]
        f0_times = times[voiced_word]
        energy_word = energy[selected]
        energy_times = times[selected]
        result[i] = (
            max(0.0, end - start), before, after, local_rate,
            float(f0.mean()) if len(f0) else 0.0,
            float(np.percentile(f0, 90) - np.percentile(f0, 10)) if len(f0) else 0.0,
            slope(f0_times, f0), float(voiced_word.sum() / max(1, selected.sum())),
            float(energy_word.mean()) if len(energy_word) else 0.0,
            float(np.percentile(energy_word, 90) - np.percentile(energy_word, 10))
            if len(energy_word) else 0.0,
            slope(energy_times, energy_word), float(zcr[selected].mean()) if selected.any() else 0.0,
        )
    return result


def read_labels(path: Path) -> list[dict]:
    sheet = load_workbook(path, read_only=True, data_only=True)["label"]
    rows = sheet.iter_rows(values_only=True)
    header = [str(value).strip() for value in next(rows)]
    columns = {name: header.index(name) for name in ("video_id", "clip_id", "text", "label", "annotation")}
    result = []
    for row in rows:
        item = {name: row[index] for name, index in columns.items()}
        clip = item["clip_id"]
        if isinstance(clip, float) and clip.is_integer():
            item["clip_id"] = str(int(clip))
        else:
            item["clip_id"] = str(clip)
        item["video_id"] = str(item["video_id"])
        result.append(item)
    if len(result) != 100:
        raise ValueError(f"Expected exactly 100 labeled clips, found {len(result)}")
    return result


def environment(text_model_name: str, face_model_path: Path, pose_model_path: Path, output: Path) -> None:
    import cv2
    import mediapipe
    import transformers
    import whisperx

    ffmpeg = subprocess.run(["ffmpeg", "-version"], capture_output=True, text=True, check=True).stdout.splitlines()[0]
    info = {"text_model": text_model_name,
            "face_model": str(face_model_path.resolve()),
            "face_model_sha256": hashlib.sha256(face_model_path.read_bytes()).hexdigest(),
            "pose_model": str(pose_model_path.resolve()),
            "pose_model_sha256": hashlib.sha256(pose_model_path.read_bytes()).hexdigest(),
            "torch": torch.__version__,
            "transformers": transformers.__version__, "mediapipe": mediapipe.__version__,
            "opencv": cv2.__version__, "whisperx": getattr(whisperx, "__version__", "record package version"),
            "ffmpeg": ffmpeg, "sample_rate_hz": SAMPLE_RATE, "audio_frame_ms": 25,
            "audio_hop_ms": 10, "video_sample_fps": 5, "face_landmark_ids": FACE_POINTS,
            "pose_landmark_ids": BODY_POINTS, "vision_frame_dim": VISION_FRAME_DIM}
    info["prosody_fields"] = PROSODY_FIELDS
    output.write_text(json.dumps(info, indent=2), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-root", type=Path, default=None)
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parent.parent / "outputs" / "features_q1_face_pose")
    parser.add_argument("--text-model", default="google-bert/bert-base-uncased")
    parser.add_argument("--face-model", type=Path, default=Path(__file__).resolve().parent.parent / "task" / "face_landmarker.task",
                        help="MediaPipe 人脸关键点模型路径")
    parser.add_argument("--pose-model", type=Path, default=Path(__file__).resolve().parent.parent / "task" / "pose_landmarker_full.task",
                        help="MediaPipe 姿态关键点模型路径")
    parser.add_argument("--device", default="auto")
    parser.add_argument("--max-samples", type=int, default=0, help="0 表示处理全部100条视频")
    args = parse_config_args(parser, "extract")
    configure_cache()
    for model_path in (args.face_model, args.pose_model):
        if not model_path.is_file():
            raise FileNotFoundError(f"Landmarker model not found: {model_path}")
    root = args.data_root or Path(__file__).resolve().parent.parent / "data"
    label_files = list(root.rglob("label-100.xlsx"))
    if len(label_files) != 1:
        raise FileNotFoundError(f"Expected one label-100.xlsx under {root}, found {len(label_files)}")
    video_root = label_files[0].parent
    labels = read_labels(label_files[0])[:args.max_samples or None]
    device = "cuda" if args.device == "auto" and torch.cuda.is_available() else (
        "cpu" if args.device == "auto" else args.device)
    if device.startswith("cuda") and not torch.cuda.is_available():
        raise RuntimeError("CUDA requested but unavailable")

    try:
        import whisperx
        from transformers import AutoModel, AutoTokenizer
    except ImportError as error:
        raise RuntimeError("Install requirements-q1.txt in the selected PyTorch environment") from error
    tokenizer = AutoTokenizer.from_pretrained(args.text_model, use_fast=True)
    text_model = AutoModel.from_pretrained(args.text_model).to(device).eval()
    if text_model.config.hidden_size != 768:
        raise ValueError("The feature export expects 768-dimensional BERT text features")
    align_model, align_metadata = whisperx.load_align_model(language_code="en", device=device)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = args.output_dir / "manifest.csv"
    manifest_fields = ["id", "video_id", "clip_id", "seconds", "words", "text_dim", "audio_dim",
                       "prosody_dim", "vision_dim", "mean_alignment_score", "alignment_text_similarity",
                       "vision_coverage", "face_coverage", "pose_coverage", "nearest_frame_fraction",
                       "status", "error"]

    with manifest_path.open("w", newline="", encoding="utf-8-sig") as manifest:
        writer = csv.DictWriter(manifest, fieldnames=manifest_fields)
        writer.writeheader()
        for item in labels:
            sample_id = f"{item['video_id']}$_${item['clip_id']}"
            video = video_root / item["video_id"] / f"{item['clip_id']}.mp4"
            row = {"id": sample_id, "video_id": item["video_id"], "clip_id": item["clip_id"],
                   "text_dim": 768, "audio_dim": 86, "prosody_dim": len(PROSODY_FIELDS),
                   "vision_dim": VISION_FRAME_DIM * 2, "status": "failed", "error": ""}
            try:
                if not video.is_file():
                    raise FileNotFoundError(video)
                audio = decode_audio(video)
                word_records = align_words(str(item["text"]), audio, align_model, align_metadata, whisperx, device)
                words = [str(value.get("word", "")).strip() for value in word_records]
                if not words:
                    raise ValueError("Forced alignment returned no word timestamps")
                text_values = text_features(words, tokenizer, text_model, device)
                at, af = audio_frames(audio)
                vt, vf = visual_frames(video, args.face_model, args.pose_model)
                audio_values, audio_mask, _ = pool_intervals(at, af, word_records)
                duration = len(audio) / SAMPLE_RATE
                prosody = prosody_features(at, af, word_records, duration)
                face_visible = vf[:, FACE_DIM - 1] > 0
                pose_visible = vf[:, -1] > 0
                visible = face_visible | pose_visible
                vision_values, vision_mask, vision_nearest = pool_intervals(
                    vt[visible], vf[visible], word_records, nearest_tolerance=0.1)
                _, face_mask, _ = pool_intervals(vt[face_visible], vf[face_visible], word_records,
                                                 nearest_tolerance=0.1)
                _, pose_mask, _ = pool_intervals(vt[pose_visible], vf[pose_visible], word_records,
                                                 nearest_tolerance=0.1)
                text_mask = np.linalg.norm(text_values, axis=1) > 0
                text_canonical = re.sub(r"[^a-z0-9']+", " ", str(item["text"]).lower()).strip()
                aligned_canonical = re.sub(r"[^a-z0-9']+", " ", " ".join(words).lower()).strip()
                similarity = SequenceMatcher(None, text_canonical, aligned_canonical).ratio()
                payload = {
                    "id": np.asarray(sample_id), "words": np.asarray(words, dtype="U"),
                    "times": np.asarray([[w["start"], w["end"]] for w in word_records], dtype=np.float32),
                    "text": text_values, "audio": audio_values, "vision": vision_values,
                    "prosody": prosody, "prosody_fields": np.asarray(PROSODY_FIELDS, dtype="U"),
                    "utterance_speech_rate_wps": np.asarray(len(words) / max(duration, 1e-6), dtype=np.float32),
                    "text_mask": text_mask.astype(np.uint8), "audio_mask": audio_mask,
                    "vision_mask": vision_mask, "face_mask": face_mask, "pose_mask": pose_mask,
                    "vision_nearest_mask": vision_nearest,
                }
                np.savez_compressed(args.output_dir / f"{item['video_id']}_{item['clip_id']}.npz", **payload)
                scores = [float(w["score"]) for w in word_records if w.get("score") is not None]
                row.update({"seconds": round(duration, 4), "words": len(words),
                            "mean_alignment_score": round(float(np.mean(scores)), 4) if scores else "",
                            "alignment_text_similarity": round(similarity, 4),
                            "vision_coverage": round(float(vision_mask.mean()), 4),
                            "face_coverage": round(float(face_mask.mean()), 4),
                            "pose_coverage": round(float(pose_mask.mean()), 4),
                            "nearest_frame_fraction": round(float(vision_nearest.mean()), 4),
                            "status": "ok"})
            except Exception as error:
                row["error"] = f"{type(error).__name__}: {error}"
                print(f"failed {sample_id}: {row['error']}", file=sys.stderr)
            writer.writerow(row)
            manifest.flush()
            print(f"{row['status']:>6} {sample_id}")

    with manifest_path.open(newline="", encoding="utf-8-sig") as manifest:
        review = [row for row in csv.DictReader(manifest)
                  if row["mean_alignment_score"]
                  and float(row["mean_alignment_score"]) < REVIEW_ALIGNMENT_SCORE]
    with (args.output_dir / "review_candidates.csv").open(
        "w", newline="", encoding="utf-8-sig"
    ) as output:
        writer = csv.DictWriter(output, fieldnames=manifest_fields)
        writer.writeheader()
        writer.writerows(review)
    environment(args.text_model, args.face_model, args.pose_model, args.output_dir / "environment.json")
    print(f"finished {len(labels)} requested rows; review {manifest_path} for failures and alignment quality")


if __name__ == "__main__":
    main()
